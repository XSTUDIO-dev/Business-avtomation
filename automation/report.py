"""Excel-отчёт и PNG-дашборд."""
import matplotlib

matplotlib.use("Agg")
import matplotlib.dates as mdates
import matplotlib.pyplot as plt
import pandas as pd
from matplotlib.patches import FancyBboxPatch
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BG, CARD, LINE, FG, MUTE = "#0b0d12", "#141824", "#242a3a", "#e8eaf0", "#8b93a7"
PALETTE = ["#4f8cff", "#9b7bff", "#34d399", "#fbbf24", "#f472b6"]


def rub(x: float) -> str:
    return f"{x:,.0f}".replace(",", "\u202f") + " ₽"


def save_excel(df, k, path):
    summary = pd.DataFrame({"Показатель": ["Выручка", "Оплаченных заказов", "Средний чек", "Доля отмен", "Повторные клиенты"],
                            "Значение": [rub(k["revenue"]), k["paid_orders"], rub(k["avg_check"]), f"{k['cancel_rate']:.1%}", f"{k['repeat_share']:.1%}"]})
    daily = k["daily"].reset_index(); daily.columns = ["Дата", "Выручка, ₽"]
    top = k["top_products"].reset_index(); top.columns = ["Товар", "Выручка, ₽"]
    with pd.ExcelWriter(path, engine="openpyxl") as xl:
        summary.to_excel(xl, sheet_name="Итоги", index=False)
        daily.to_excel(xl, sheet_name="По дням", index=False)
        top.to_excel(xl, sheet_name="Топ товаров", index=False)
        df.to_excel(xl, sheet_name="Чистые данные", index=False)
        for ws in xl.book.worksheets:
            for c in ws[1]:
                c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="2B3A67"); c.alignment = Alignment(horizontal="center")
            for i, col in enumerate(ws.columns, 1):
                ws.column_dimensions[get_column_letter(i)].width = min(46, max(len(str(c.value or "")) for c in col) + 3)
            ws.freeze_panes = "A2"
            ws.page_setup.orientation, ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = "landscape", 1, 0
            ws.sheet_properties.pageSetUpPr.fitToPage = True


def _style(ax, title):
    ax.set_facecolor(CARD)
    for s in ax.spines.values():
        s.set_color(LINE)
    ax.grid(color=LINE, lw=.6, alpha=.8); ax.set_axisbelow(True); ax.tick_params(labelsize=9)
    ax.set_title(title, loc="left", color=MUTE, fontsize=11, pad=10)


def save_dashboard(k, path, title):
    plt.rcParams.update({"font.family": "DejaVu Sans", "text.color": FG, "xtick.color": MUTE, "ytick.color": MUTE})
    fig = plt.figure(figsize=(13, 7.4), facecolor=BG)
    gs = fig.add_gridspec(3, 4, height_ratios=[.5, 1.5, 1.4], hspace=.62, wspace=.34, left=.05, right=.97, top=.9, bottom=.07)
    fig.text(.05, .945, title, fontsize=18, fontweight="bold")
    cards = [("Выручка", rub(k["revenue"])), ("Оплаченных заказов", str(k["paid_orders"])),
             ("Средний чек", rub(k["avg_check"])), ("Повторные клиенты", f"{k['repeat_share']:.0%}")]
    for i, (name, val) in enumerate(cards):
        ax = fig.add_subplot(gs[0, i]); ax.axis("off")
        ax.add_patch(FancyBboxPatch((0, 0), 1, 1, boxstyle="round,pad=0,rounding_size=0.12", transform=ax.transAxes, fc=CARD, ec=LINE))
        ax.text(.08, .6, name, color=MUTE, fontsize=10.5, transform=ax.transAxes)
        ax.text(.08, .17, val, fontsize=20, fontweight="bold", transform=ax.transAxes)
    ax = fig.add_subplot(gs[1, :3]); _style(ax, "Выручка по дням")
    d = k["daily"]
    ax.plot(d.index, d.values, color=PALETTE[0], lw=2.2); ax.fill_between(d.index, d.values, color=PALETTE[0], alpha=.15)
    ax.xaxis.set_major_formatter(mdates.DateFormatter("%d.%m"))
    ax = fig.add_subplot(gs[1, 3]); ax.set_facecolor(CARD); ax.set_title("Каналы продаж", loc="left", color=MUTE, fontsize=11, pad=10)
    ch = k["by_channel"]
    w, _ = ax.pie(ch.values, colors=PALETTE, startangle=90, wedgeprops=dict(width=.38, edgecolor=CARD))
    ax.legend(w, [f"{n} · {v / ch.sum():.0%}" for n, v in ch.items()], loc="upper center", bbox_to_anchor=(.5, .02),
              frameon=False, labelcolor=FG, fontsize=8.5, ncol=2)
    ax = fig.add_subplot(gs[2, :3]); _style(ax, "Топ-5 товаров по выручке")
    t = k["top_products"][::-1]
    ax.barh(t.index, t.values, color=PALETTE[1], height=.55)
    for y, v in enumerate(t.values):
        ax.text(v, y, "  " + rub(v), va="center", fontsize=9, color=FG)
    ax.set_xlim(0, t.max() * 1.18)
    ax = fig.add_subplot(gs[2, 3]); _style(ax, "Статусы заказов")
    s = k["status"]
    ax.bar(s.index, s.values, color=[PALETTE[2], PALETTE[3], PALETTE[4]][:len(s)], width=.55)
    fig.savefig(path, dpi=130, facecolor=BG)
    plt.close(fig)
